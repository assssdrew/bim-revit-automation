import openpyxl,csv,sys
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L
src=sys.argv[1]; out=sys.argv[2]; marker=sys.argv[3] if len(sys.argv)>3 else ''
if src=='empty':
    rows=[['Этаж','Тип','Номер','Sобщ','Sж','Имя','Число']]
elif src.endswith('.txt'):
    rows=list(csv.reader(open(src,encoding='utf-8-sig'),delimiter=';'))
else:
    ws0=openpyxl.load_workbook(src)['Спецификация СУ+УБ+ГР 2.1']
    rows=[[('' if v is None else str(v)) for v in r] for r in ws0.iter_rows(values_only=True)]
N=3001; NA=1500; NF=80
wb=openpyxl.Workbook(); B=Font(bold=True)
wv=wb.active; wv.title='Выгрузка'
for r in rows: wv.append(r)
if marker: wv.append(['Связь АР:',marker]); wv.append(['Выгружено:','19.10.2026 12:30'])
# Настройки
s=wb.create_sheet('Настройки')
s['A1']='НАСТРОЙКИ (меняются без правки формул)';s['A1'].font=Font(bold=True,size=13)
s['A3']='Корпус (формула берёт K__ из имени связи АР; можно вписать вручную)'
s['E2']='служебное';s['D3']='Связь АР:'
s['E3']='=IFERROR(INDEX(INDIRECT("\'Выгрузка\'!B1:B4000"),MATCH("Связь АР:",INDIRECT("\'Выгрузка\'!A1:A4000"),0))&"","")'
RR='ROW($1:$80)'
s['E4']=('=IFERROR(LOOKUP(2,1/(((MID($E$3,'+RR+',1)="K")+(MID($E$3,'+RR+',1)="К"))*(MID("_"&$E$3,'+RR+',1)="_")'
        '*ISNUMBER(--MID($E$3,'+RR+'+1,2))*(LEN(MID($E$3,'+RR+'+1,2))=2)),'+RR+'),"")')
s['B3']='=IF(E4<>"","K"&MID(E3,E4+1,2),IF(E3<>"",E3,"—"))'
s.column_dimensions['E'].width=34
s['D5']='Выгружено:'
s['E5']='=IFERROR(INDEX(INDIRECT("\'Выгрузка\'!B1:B4000"),MATCH("Выгружено:",INDIRECT("\'Выгрузка\'!A1:A4000"),0))&"","")'
s['A4']='Учитывать тип квартиры при поиске типовых этажей (да/нет)';s['B4']='нет'
s['A6']='Номер столбца в «Выгрузке» (A=1, B=2 …)';s['A6'].font=B
for i,(n,v) in enumerate([('Этаж',1),('Тип квартиры',2),('Номер квартиры на этаже',3),('Площадь общая',4),('Площадь жилая',5),('Имя помещения',6),('Число (пусто = 1)',7)]):
    s.cell(7+i,1,n);s.cell(7+i,2,v)
s['A15']='Имена помещений (по одному в ячейке, регистр не важен)';s['A15'].font=B
for j,(h,vals) in enumerate([('СУ',['Совмещенный санузел','Санузел','Ванная','Совмещённый санузел']),('Уборная',['Уборная','Туалет']),('Гардеробная',['Гардеробная'])]):
    c=2+j;s.cell(16,c,h).font=B
    for k,v in enumerate(vals):s.cell(17+k,c,v)
s['A17']='варианты →';s['A27']='Строки 17–26 — до 10 вариантов на каждый вид. Тип «1(1XL)» разбирается: Комнат=1, Тип=1XL.'
s.column_dimensions['A'].width=58;
for c in 'BCD':s.column_dimensions[c].width=24
SU="Настройки!$B$17:$B$26";UB="Настройки!$C$17:$C$26";GR="Настройки!$D$17:$D$26"
dec='MID(1/2,2,1)'
def num(x): return f'IFERROR(--SUBSTITUTE(SUBSTITUTE(TRIM({x}),",","."),".",{dec}),"")'
# Расчет (helper)
h=wb.create_sheet('Расчет')
for c,t in enumerate(['Этаж','Номер','Тип','Sобщ','Sж','Имя','Число','Вид','Ключ кв.','№ кв. п/п'],1):h.cell(1,c,t).font=B
def g(k,r): return f'INDEX(INDIRECT("\'Выгрузка\'!A{r}:Z{r}"),Настройки!$B${6+k})'
for r in range(2,N+1):
    h[f'A{r}']=f'={num(g(1,r))}'
    h[f'B{r}']=f'={num(g(3,r))}'
    h[f'C{r}']=f'=TRIM({g(2,r)}&"")'
    h[f'D{r}']=f'={num(g(4,r))}'
    h[f'E{r}']=f'={num(g(5,r))}'
    h[f'F{r}']=f'=TRIM({g(6,r)}&"")'
    h[f'G{r}']=f'=IF(F{r}="",0,IF(TRIM({g(7,r)}&"")="",1,{num(g(7,r))}))'
    h[f'H{r}']=f'=IF(F{r}="",0,IF(COUNTIF({SU},F{r})>0,1,IF(COUNTIF({UB},F{r})>0,2,IF(COUNTIF({GR},F{r})>0,3,0))))'
    h[f'I{r}']=f'=IF(AND(ISNUMBER(A{r}),ISNUMBER(B{r})),A{r}*1000+B{r},"")'
    h[f'J{r}']=f'=IF(I{r}="","",IF(COUNTIF(I$2:I{r},I{r})=1,MAX(J$1:J{r-1})+1,""))'
R=lambda c:f'Расчет!${c}$2:${c}${N}'
# Квартиры
q=wb.create_sheet('Квартиры')
hd=['ID','Этаж','Номер','Тип','Sобщ','Sж','Комнаты','СУ','Уборная','Гардеробная','Ключ','строка']
for c,t in enumerate(hd,1):q.cell(1,c,t).font=B
for r in range(2,NA+2):
    m=f'$L{r}'
    q[f'L{r}']=f'=IFERROR(MATCH(ROW()-1,{R("J")},0),"")'
    q[f'K{r}']=f'=IF({m}="","",INDEX({R("I")},{m}))'
    q[f'B{r}']=f'=IF({m}="","",INDEX({R("A")},{m}))'
    q[f'C{r}']=f'=IF({m}="","",INDEX({R("B")},{m}))'
    t=f'INDEX({R("C")},{m})'
    q[f'D{r}']=f'=IF({m}="","",IF(ISNUMBER(FIND("(",{t})),SUBSTITUTE(MID({t},FIND("(",{t})+1,99),")",""),{t}))'
    q[f'E{r}']=f'=IF({m}="","",INDEX({R("D")},{m}))'
    q[f'F{r}']=f'=IF({m}="","",INDEX({R("E")},{m}))'
    q[f'G{r}']=f'=IF({m}="","",IFERROR(--LEFT({t},FIND("(",{t})-1),IFERROR(--LEFT({t},1),"")))'
    for col,k in (('H',1),('I',2),('J',3)):
        q[f'{col}{r}']=f'=IF({m}="","",SUMIFS({R("G")},{R("I")},$K{r},{R("H")},{k}))'
    q[f'A{r}']=f'=IF({m}="","",B{r}&"."&C{r}&"-"&D{r}&"-"&H{r}&"."&I{r}&"."&J{r})'
q.freeze_panes='B2'; q.column_dimensions['A'].width=18
Q=lambda c:f'Квартиры!${c}$2:${c}${NA+1}'
# Типовые этажи
t=wb.create_sheet('Типовые этажи',0)
t['A1']='=Настройки!B3';t['A1'].font=Font(bold=True,size=13)
t['C1']='=IF(LEFT(INDIRECT("\'Выгрузка\'!A1")&"",6)="ОШИБКА","Ошибка выгрузки — см. лист Выгрузка","Квартир: "&COUNT('+Q('K')+')&"   Этажей: "&COUNT(A3:A'+str(NF+2)+')&"   Групп: "&MAX(B3:B'+str(NF+2)+')&IF(Настройки!B4="да","   (с учётом типа)",""))'
hd=['Этаж','Группа']+[f'Кв. {i}' for i in range(1,16)]+['Ключ','Этажи группы','','Группа','Этажи','Кол-во']
for c,x in enumerate(hd,1):t.cell(2,c,x).font=B
last=NF+2
for r in range(3,last+1):
    t[f'A{r}']=(f'=IFERROR(MIN({Q("B")})/(COUNT({Q("B")})>0),"")' if r==3 else
               f'=IFERROR(IF(A{r-1}="","",SMALL({Q("B")},COUNTIF({Q("B")},"<="&A{r-1})+1)),"")')
    for i in range(1,16):
        c=L(2+i); m=f'MATCH($A{r}*1000+{i},{Q("K")},0)'
        ix=lambda col:f'INDEX({Q(col)},{m})'
        t[f'{c}{r}']=(f'=IF($A{r}="","",IFERROR(IF(Настройки!$B$4="да",{ix("D")}&" / ","")&{ix("H")}&" СУ"'
                      f'&IF({ix("I")}>0," +У","")&IF({ix("J")}>0," +Г",""),"-"))')
    t[f'R{r}']='=IF(A%d="","",'%r+'&"|"&'.join(f'{L(2+i)}{r}' for i in range(1,16))+')'
    t[f'B{r}']=f'=IF(A{r}="","",IF(COUNTIF(R$2:R{r-1},R{r})=0,MAX(B$2:B{r-1})+1,INDEX(B$2:B{r-1},MATCH(R{r},R$2:R{r-1},0))))'
    t[f'S{r}']=f'=IF(A{r}="","",IFERROR(LOOKUP(2,1/(B$2:B{r-1}=B{r}),S$2:S{r-1})&", ","")&A{r})'
    k=r-2
    t[f'U{r}']=f'=IF({k}>MAX(B$3:B${last}),"",{k})'
    t[f'V{r}']=f'=IF(U{r}="","",LOOKUP(2,1/(B$3:B${last}=U{r}),S$3:S${last}))'
    t[f'W{r}']=f'=IF(U{r}="","",COUNTIF(B$3:B${last},U{r}))'
cols=['FFC7CE','C6EFCE','FFEB9C','BDD7EE','E4DFEC','FCE4D6','D9D9D9','DDEBF7','E2EFDA','FFF2CC','F8CBAD','C9C9FF']
for i,cl in enumerate(cols):
    f=PatternFill(start_color=cl,end_color=cl,fill_type='solid')
    t.conditional_formatting.add(f'A3:Q{last}',FormulaRule(formula=[f'AND($B3<>"",MOD($B3-1,12)={i})'],fill=f))
    t.conditional_formatting.add(f'U3:W{last}',FormulaRule(formula=[f'AND($U3<>"",MOD($U3-1,12)={i})'],fill=f))
t.column_dimensions['R'].hidden=True
t['N1']='=IF(Настройки!E5="","","Выгружено: "&Настройки!E5)';t['N1'].font=Font(italic=True)
for i in range(3,18):t.column_dimensions[L(i)].width=13
t.column_dimensions['S'].width=30;t.column_dimensions['V'].width=40
t.freeze_panes='C3'

# Таблица 1
from openpyxl.styles import Border,Side
T=wb.create_sheet('Таблица 1',1)
T.merge_cells('A1:J1');T['A1']='="Корпус "&Настройки!B3';T['A1'].font=Font(bold=True,size=12)
T['A1'].fill=PatternFill(start_color='92D050',end_color='92D050',fill_type='solid');T['A1'].alignment=Alignment(horizontal='center')
for c,x in enumerate(['ID','Этаж','Номер','Тип','Комнат','Sобщ','Sж','СУ','Уборная','Гардеробная'],1):
    T.cell(2,c,x).font=B;T.cell(3,c,c).alignment=Alignment(horizontal='center')
mp='ABCDGEFHIJ'
for r in range(4,NA+4):
    for c,src in enumerate(mp,1):
        T.cell(r,c,f'=IF(Квартиры!{src}{r-2}="","",Квартиры!{src}{r-2})')
T.freeze_panes='B4';T.column_dimensions['A'].width=18

I=wb.create_sheet('Инструкция',0)
steps=['Выгрузка квартирографии — как получить отчёт','',
'Один раз: в Проигрывателе Dynamo выберите папку \\\\server\\share\\BIM\\Dynamo\\Выгрузка квартирографии.','',
'Каждый раз:',
'1. Закройте файлы Квартирография_*.xlsx.',
'2. Откройте модель ОВ нужного корпуса (связь АР загружена, модель сохранена).',
'3. Управление → Проигрыватель Dynamo → «Выгрузка_квартирографии» → Запуск. «Связь АР» оставьте пустым.',
'4. Откройте папку «Квартирография» рядом с моделью ОВ (для совместной работы — рядом с файлом-хранилищем) и файл Квартирография_K__.xlsx. Результат: листы «Типовые этажи» и «Таблица 1» (F9 при необходимости).','',
'Если ошибка:',
'• Появился файл Квартирография_ОШИБКА.xlsx — в нём на листе «Выгрузка» причина и список связей. Загрузите связь АР или впишите часть её имени в поле «Связь АР» (если связей АР несколько).',
'• Узел жёлтый с текстом «Закройте файл…» — закройте Excel и запустите снова. «Нет шаблона…» — верните Квартирография_ШАБЛОН.xlsx в папку скрипта. «Модель ОВ не сохранена / в облаке» — сохраните модель на сервере.',
'• «ВНИМАНИЕ: нет параметров…» — поправьте имена параметров в блоке НАСТРОЙКИ узла «Выгрузка АР». Имена помещений — на листе «Настройки».']
for i,x in enumerate(steps,1): I.cell(i,1,x)
I['A1'].font=Font(bold=True,size=14);I['A5'].font=B;I['A11'].font=B;I.column_dimensions['A'].width=120
wb.move_sheet('Типовые этажи',1-wb.sheetnames.index('Типовые этажи'))
order=['Инструкция','Типовые этажи','Таблица 1','Выгрузка','Настройки','Расчет','Квартиры']
wb._sheets=[wb[n] for n in order]
wb['Расчет'].sheet_state='hidden';wb['Квартиры'].sheet_state='hidden'
wb.active=0
wb.save(out)
